from dataclasses import dataclass
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from docx import Document as WordDocument
from pypdf import PdfReader


MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_PDF_PAGES = 300
MAX_EXTRACTED_CHARACTERS = 1_000_000
MAX_DOCX_ENTRIES = 2048
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024


@dataclass(frozen=True)
class ExtractionResult:
    status: str
    text: str | None
    reason: str | None = None


def _clean_text(value: str) -> str:
    """Remove null and disallowed control characters."""
    return "".join(
        character
        for character in value
        if character in "\n\r\t" or ord(character) >= 32
    ).strip()


def _validate_docx_archive(content: bytes) -> str | None:
    """Validate DOCX archive structure and basic resource limits."""
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()

            if len(entries) > MAX_DOCX_ENTRIES:
                return "The Word file contains too many archive entries."

            total_size = sum(entry.file_size for entry in entries)

            if total_size > MAX_DOCX_UNCOMPRESSED_BYTES:
                return "The Word file exceeds the extraction size limit."

            if any(
                entry.file_size > MAX_DOCX_UNCOMPRESSED_BYTES
                for entry in entries
            ):
                return "A Word file component exceeds the extraction size limit."

            required_entries = {
                "[Content_Types].xml",
                "word/document.xml",
            }
            available_entries = {entry.filename for entry in entries}

            if not required_entries.issubset(available_entries):
                return "The file is not a valid DOCX document."

    except (BadZipFile, OSError, ValueError):
        return "The Word document archive is invalid or incomplete."

    return None


def _extract_pdf(content: bytes) -> ExtractionResult:
    """Extract text without modifying the original PDF."""
    if b"%PDF-" not in content[:1024]:
        return ExtractionResult(
            "Failed", None, "The file does not contain a valid PDF signature."
        )

    try:
        reader = PdfReader(BytesIO(content), strict=False)

        if reader.is_encrypted:
            return ExtractionResult(
                "Failed",
                None,
                "Password-protected PDFs cannot currently be indexed.",
            )

        if len(reader.pages) > MAX_PDF_PAGES:
            return ExtractionResult(
                "Unsupported", None, "The PDF exceeds the configured page limit."
            )

        pages = []
        character_count = 0

        for page in reader.pages:
            page_text = page.extract_text() or ""
            character_count += len(page_text)

            if character_count > MAX_EXTRACTED_CHARACTERS:
                return ExtractionResult(
                    "Unsupported",
                    None,
                    "Extracted text exceeds the configured size limit.",
                )

            pages.append(page_text)

        extracted_text = _clean_text("\n\n".join(pages))

    except Exception:
        return ExtractionResult(
            "Failed", None, "The PDF could not be read for text extraction."
        )

    if not extracted_text:
        return ExtractionResult(
            "NeedsOCR",
            None,
            "No selectable text was found; OCR may be required.",
        )

    return ExtractionResult("Extracted", extracted_text)


def _extract_docx(content: bytes) -> ExtractionResult:
    """Extract text from DOCX paragraphs, tables, headers and footers."""
    archive_error = _validate_docx_archive(content)

    if archive_error:
        return ExtractionResult("Failed", None, archive_error)

    try:
        document = WordDocument(BytesIO(content))
        parts = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)

        for table in document.tables:
            for row in table.rows:
                cell_text = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]
                if cell_text:
                    parts.append(" | ".join(cell_text))

        for section in document.sections:
            for container in (section.header, section.footer):
                for paragraph in container.paragraphs:
                    if paragraph.text.strip():
                        parts.append(paragraph.text)

        combined_text = _clean_text("\n".join(parts))

    except Exception:
        return ExtractionResult(
            "Failed",
            None,
            "The Word document could not be read for text extraction.",
        )

    if len(combined_text) > MAX_EXTRACTED_CHARACTERS:
        return ExtractionResult(
            "Unsupported",
            None,
            "Extracted text exceeds the configured size limit.",
        )

    if not combined_text:
        return ExtractionResult(
            "Extracted", "", "The Word document contains no extractable text."
        )

    return ExtractionResult("Extracted", combined_text)


def extract_document_text(
    filename: str,
    file_content: bytes,
) -> ExtractionResult:
    """Extract supported document text without altering the original file."""
    if not file_content:
        return ExtractionResult("Failed", None, "The file is empty.")

    if len(file_content) > MAX_FILE_BYTES:
        return ExtractionResult(
            "Unsupported", None, "The file exceeds the configured size limit."
        )

    normalized_filename = (filename or "").lower().strip()

    if normalized_filename.endswith(".pdf"):
        return _extract_pdf(file_content)

    if normalized_filename.endswith(".docx"):
        return _extract_docx(file_content)

    return ExtractionResult(
        "Unsupported",
        None,
        "Content indexing currently supports PDF and DOCX files only.",
    )
